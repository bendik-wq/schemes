import random
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
random.seed(7)
F='Arial'
BLUE=Font(name=F,color='0000FF',size=10); BLK=Font(name=F,size=10); GRN=Font(name=F,color='008000',size=10)
BOLD=Font(name=F,bold=True,size=10); H1=Font(name=F,bold=True,size=14); IT=Font(name=F,italic=True,size=9,color='666666')
HDR=Font(name=F,bold=True,size=10,color='FFFFFF'); HFILL=PatternFill('solid',fgColor='1F3864')
YEL=PatternFill('solid',fgColor='FFFF00'); SEC=PatternFill('solid',fgColor='D9E1F2'); RED=PatternFill('solid',fgColor='FCE4D6')
TOP=Border(top=Side(style='thin')); DBL=Border(top=Side(style='thin'),bottom=Side(style='double'))
USD='$#,##0;($#,##0);"-"'; PCT='0.0%;(0.0%);"-"'; X='0.00"x"'; DAYS='0" days"'
def n(k): # realistic non-round dollars from $k
    return int(round(k*1000*(1+random.uniform(-0.006,0.006))))
wb=Workbook()
def sheet(name,title,sub,cols):
    ws=wb.create_sheet(name); ws['A1']=title; ws['A1'].font=H1
    ws['A2']=sub; ws['A2'].font=IT
    ws['A3']='Kestrel Plumbing & Gas Pty Ltd is a fictional company built for teaching. Any resemblance to a real business is coincidental.'; ws['A3'].font=IT
    ws.column_dimensions['A'].width=46
    for i,c in enumerate(cols):
        cell=ws.cell(5,2+i,c); cell.font=HDR; cell.fill=HFILL; cell.alignment=Alignment(horizontal='center')
        ws.column_dimensions[chr(66+i)].width=15
    ws.cell(5,1,'A$' if cols else '').font=HDR; ws.cell(5,1).fill=HFILL
    ws.sheet_view.showGridLines=False; ws.freeze_panes='B6'
    return ws
def put(ws,r,label,vals,font=BLK,fmt=USD,bold=False,border=None,fill=None,note=None,start=2):
    ws.cell(r,1,label).font=BOLD if bold else BLK
    for i,v in enumerate(vals):
        c=ws.cell(r,start+i,v); c.number_format=fmt
        c.font=Font(name=F,size=10,bold=bold,color=font.color.rgb if font.color else None)
        if border: c.border=border
        if fill: c.fill=fill
    if fill: ws.cell(r,1).fill=fill
    if note:
        col=start+len(vals); ws.cell(r,col,note).font=IT
def section(ws,r,label,ncol):
    for c in range(1,ncol+2): ws.cell(r,c).fill=SEC
    ws.cell(r,1,label).font=BOLD
wb.remove(wb.active)

# ---------------- README
rd=wb.create_sheet('Read me'); rd.sheet_view.showGridLines=False; rd.column_dimensions['A'].width=30; rd.column_dimensions['B'].width=95
rows=[('Kestrel Plumbing & Gas Pty Ltd',None),('Example acquisition target for technical videos. Fictional.',None),('',None),
('The business','Commercial and residential plumbing and gas. 24 staff (18 field). One depot, 14 vans. Owner-founder, 61, wants to retire. Financial year ends 30 June. Figures in A$, ex GST.'),
('Story in the numbers','Revenue up 22% in two years, but gross margin is sliding, debtor days are stretching, one builder is 21% of revenue, and the owner still works on the tools two days a week. The seller says adjusted EBITDA is about $968K. A buyer gets to about $614K.'),
('',None),('Tabs',None),
('P&L','Three years of statutory-style profit and loss, revenue by service line.'),
('Balance Sheet','Four year-ends (FY22 opening). Cash is the balancing item; drivers (debtor days etc.) sit at the bottom.'),
('Cash Flow','Indirect method, built from the P&L and balance-sheet movements. Reconciles to cash.'),
('Normalisation','EBITDA add-backs: what the seller claims vs what a buyer accepts, line by line.'),
('Working Capital','Debtor, creditor, WIP and stock days, NWC, the NWC peg and the completion adjustment.'),
('Valuation Bridge','Enterprise value to equity: debt and debt-like items, surplus cash, NWC adjustment.'),
('Funding & DSCR','Bank + vendor finance, repayments, CFADS and DSCR at the buyer price vs the seller asking price.'),
('Red Flags','The diligence checks that move the price, each calculated from the statements.'),
('Video ideas','Technical video ideas, each tied to a tab.'),
('',None),('Colour code',None),
('Blue text','Hard-coded input. Change it and everything recalculates.'),
('Black text','Formula.'),('Green text','Pulled from another tab.'),('Yellow fill','Key assumption to play with on camera.'),
('Red fill','A check or flag.')]
for i,(a,b) in enumerate(rows,1):
    rd.cell(i,1,a).font=H1 if i==1 else (BOLD if b is not None or a in('Tabs','Colour code') else IT)
    if b: rd.cell(i,2,b).font=BLK; rd.cell(i,2).alignment=Alignment(wrap_text=True,vertical='top')
rd.cell(19,1).font=Font(name=F,color='0000FF',bold=True); rd.cell(21,1).font=Font(name=F,color='008000',bold=True)
rd.cell(22,1).fill=YEL; rd.cell(23,1).fill=RED

# ---------------- P&L  (cols B,C,D = FY23,FY24,FY25)
pl=sheet('P&L','Kestrel Plumbing & Gas Pty Ltd','Profit and loss statement · years ended 30 June · A$ ex GST',['FY23','FY24','FY25'])
R={}; r=6
def line(key,label,ks,note=None):
    global r; put(pl,r,label,[n(k) for k in ks],font=BLUE,note=note); R[key]=r; r+=1
def tot(key,label,formula,bold=True,border=TOP,fmt=USD,note=None,fill=None):
    global r; put(pl,r,label,[formula.format(c=c) for c in 'BCD'],bold=bold,border=border,fmt=fmt,note=note,fill=fill); R[key]=r; r+=1
section(pl,r,'Revenue',4); r+=1
line('rm','Maintenance contracts',[1150,1280,1410],'Annual service agreements, strata and commercial')
line('rr','Reactive and emergency',[1320,1410,1520],'Call-outs, after-hours')
line('rp','Projects and new build',[980,1150,1270],'Includes one builder, see Red Flags')
tot('rev','Total revenue','=SUM({c}%d:{c}%d)'%(R['rm'],R['rp']))
r+=1; section(pl,r,'Cost of sales',4); r+=1
line('cl','Direct labour (field staff incl. super)',[1180,1300,1420])
line('cm','Materials',[720,810,890])
line('cs','Subcontractors',[190,230,280])
line('cf','Fleet running costs (fuel, rego, servicing)',[120,132,141])
tot('cogs','Total cost of sales','=SUM({c}%d:{c}%d)'%(R['cl'],R['cf']))
r+=1; tot('gp','Gross profit','={c}%d-{c}%d'%(R['rev'],R['cogs']))
tot('gm','Gross margin %','=IF({c}%d=0,0,{c}%d/{c}%d)'%(R['rev'],R['gp'],R['rev']),bold=False,border=None,fmt=PCT)
r+=1; section(pl,r,'Operating expenses',4); r+=1
line('os','Director salary (owner, incl. super)',[180,180,180],'Owner also works on the tools ~2 days/week')
line('ow','Spouse wages (bookkeeping)',[65,68,70],'Family member on payroll')
line('oa','Office and admin wages',[210,225,240])
line('ore','Rent: depot (paid to owner\'s property trust)',[96,96,96],'Related-party lease')
line('omk','Marketing and advertising',[38,45,52])
line('oi','Insurance',[41,46,50])
line('oit','Software and IT',[22,26,31])
line('ol','Accounting and legal',[28,30,62],'FY25 includes contract dispute with a builder')
line('ov','Motor vehicles: owner and spouse cars',[34,36,38])
line('ou','Telephone and utilities',[19,20,21])
line('ot','Training and licences',[12,14,15])
line('ob','Bad debts written off',[6,9,28])
line('oo','Other expenses',[25,27,29])
tot('opex','Total operating expenses','=SUM({c}%d:{c}%d)'%(R['os'],R['oo']))
r+=1; tot('ebitda','EBITDA (reported)','={c}%d-{c}%d'%(R['gp'],R['opex']),fill=YEL)
tot('em','EBITDA margin %','=IF({c}%d=0,0,{c}%d/{c}%d)'%(R['rev'],R['ebitda'],R['rev']),bold=False,border=None,fmt=PCT)
line('da','Depreciation and amortisation',[58,66,74])
tot('ebit','EBIT','={c}%d-{c}%d'%(R['ebitda'],R['da']))
line('int','Interest (equipment finance)',[14,17,19])
tot('pbt','Net profit before tax','={c}%d-{c}%d'%(R['ebit'],R['int']))
pl.cell(r,1,'Company tax rate').font=BLK; pl.cell(r,2,0.25).font=BLUE; pl.cell(r,2).number_format=PCT; pl.cell(r,3,'Base rate entity').font=IT; R['tr']=r; r+=1
tot('tax','Income tax expense','=ROUND({c}%d*$B$%d,0)'%(R['pbt'],R['tr']),bold=False,border=None)
tot('npat','Net profit after tax','={c}%d-{c}%d'%(R['pbt'],R['tax']),border=DBL)
PL=R

# ---------------- Balance sheet (B..E = FY22..FY25)
bs=sheet('Balance Sheet','Kestrel Plumbing & Gas Pty Ltd','Balance sheet · as at 30 June · A$',['FY22','FY23','FY24','FY25'])
B={}; r=6
plc={'C':'B','D':'C','E':'D'}  # BS col -> P&L col
def bl(key,label,ks,note=None):
    global r; put(bs,r,label,[n(k) for k in ks],font=BLUE,note=note); B[key]=r; r+=1
def bf(key,label,first,formula,bold=False,border=None,font=BLK,note=None,fill=None,fmt=USD):
    """first = FY22 value (number or formula); formula uses {c} (this col), {p} (prior col), {l} (P&L col)"""
    global r
    vals=[first]+[formula.format(c=c,p=chr(ord(c)-1),l=plc[c]) for c in 'CDE']
    put(bs,r,label,vals,bold=bold,border=border,note=note,fill=fill,fmt=fmt)
    if isinstance(first,(int,float)): bs.cell(r,2).font=Font(name=F,size=10,color='0000FF',bold=bold)
    for c in 'CDE':
        if "'P&L'" in formula: bs[f'{c}{r}'].font=Font(name=F,size=10,color='008000',bold=bold)
    B[key]=r; r+=1
# placeholders for drivers rows (filled later); we need their row numbers, so define drivers block position up front
DRV=50
d={'dd':DRV+1,'cd':DRV+2,'capex':DRV+3,'div':DRV+4}
section(bs,r,'Current assets',5); r+=1
B['cash']=r; r+=1  # filled after totals known
bf('deb','Trade debtors',n(520),"=ROUND('P&L'!{l}%d*{c}%d/365,0)"%(PL['rev'],d['dd']),note='Driven by debtor days')
bl('wip','Contract assets (unbilled work in progress)',[41,48,62,85])
bl('inv','Inventory (van and depot stock)',[57,62,68,74])
bl('pre','Prepayments',[16,18,20,22])
bl('dl','Loan to director',[60,85,110,140],'Div 7A loan. Not an operating asset')
bf('tca','Total current assets','=SUM(B%d:B%d)'%(B['cash'],B['dl']),'=SUM({c}%d:{c}%d)'%(B['cash'],B['dl']),bold=True,border=TOP)
r+=1; section(bs,r,'Non-current assets',5); r+=1
bf('ppe','Plant, equipment and vehicles at cost',n(460),'={p}%d+{c}%d'%(0,d['capex']))
# fix ppe formula (prior row same)
for c in 'CDE': bs[f'{c}{B["ppe"]}']=f'={chr(ord(c)-1)}{B["ppe"]}+{c}{d["capex"]}'
bf('acc','Less accumulated depreciation',-n(152),"={p}%d-'P&L'!{l}%d"%(0,PL['da']))
for c in 'CDE': bs[f'{c}{B["acc"]}']=f"={chr(ord(c)-1)}{B['acc']}-'P&L'!{plc[c]}{PL['da']}"
bf('nppe','Net plant, equipment and vehicles','=B%d+B%d'%(B['ppe'],B['acc']),'={c}%d+{c}%d'%(B['ppe'],B['acc']),bold=True,border=TOP)
r+=1; bf('ta','TOTAL ASSETS','=B%d+B%d'%(B['tca'],B['nppe']),'={c}%d+{c}%d'%(B['tca'],B['nppe']),bold=True,border=DBL)
r+=1; section(bs,r,'Current liabilities',5); r+=1
bf('cre','Trade creditors',n(98),"=ROUND(('P&L'!{l}%d+'P&L'!{l}%d)*{c}%d/365,0)"%(PL['cm'],PL['cs'],d['cd']),note='Materials + subcontractors x creditor days')
bl('accr','Accrued expenses',[39,42,46,51])
bl('gst','GST payable',[52,58,64,71])
bl('payg','PAYG withholding and super payable',[43,47,52,57])
bl('al','Provision for annual leave',[88,96,108,121],'Owed to staff')
bl('lslc','Provision for long service leave (current)',[29,34,40,48])
bl('itp','Income tax payable',[27,30,38,41])
bl('efc','Equipment finance (current)',[38,42,48,55])
bf('tcl','Total current liabilities','=SUM(B%d:B%d)'%(B['cre'],B['efc']),'=SUM({c}%d:{c}%d)'%(B['cre'],B['efc']),bold=True,border=TOP)
r+=1; section(bs,r,'Non-current liabilities',5); r+=1
bl('lsln','Provision for long service leave (non-current)',[47,52,58,63])
bl('efn','Equipment finance (non-current)',[104,118,142,160])
bf('tncl','Total non-current liabilities','=SUM(B%d:B%d)'%(B['lsln'],B['efn']),'=SUM({c}%d:{c}%d)'%(B['lsln'],B['efn']),bold=True,border=TOP)
r+=1; bf('tl','TOTAL LIABILITIES','=B%d+B%d'%(B['tcl'],B['tncl']),'={c}%d+{c}%d'%(B['tcl'],B['tncl']),bold=True,border=TOP)
r+=1; bf('na','NET ASSETS','=B%d-B%d'%(B['ta'],B['tl']),'={c}%d-{c}%d'%(B['ta'],B['tl']),bold=True,border=DBL)
r+=1; section(bs,r,'Equity',5); r+=1
put(bs,r,'Share capital',[100,100,100,100],font=BLUE,note='100 ordinary shares at $1'); B['sc']=r; r+=1
bf('re','Retained earnings',n(520),"={p}%d+'P&L'!{l}%d-{c}%d"%(0,PL['npat'],d['div']))
for c in 'CDE': bs[f'{c}{B["re"]}']=f"={chr(ord(c)-1)}{B['re']}+'P&L'!{plc[c]}{PL['npat']}-{c}{d['div']}"
bf('te','TOTAL EQUITY','=B%d+B%d'%(B['sc'],B['re']),'={c}%d+{c}%d'%(B['sc'],B['re']),bold=True,border=DBL)
r+=1; bf('chk','Check: net assets less equity (must be 0)','=ROUND(B%d-B%d,0)'%(B['na'],B['te']),'=ROUND({c}%d-{c}%d,0)'%(B['na'],B['te']),fill=RED)
# cash = balancing item: total liabilities + equity - other assets
other='+'.join(['{c}%d'%B[k] for k in ['deb','wip','inv','pre','dl','nppe']])
put(bs,B['cash'],'Cash at bank',[f'=B{B["tl"]}+B{B["te"]}-('+other.format(c='B')+')']+[f'={c}{B["tl"]}+{c}{B["te"]}-('+other.format(c=c)+')' for c in 'CDE'],note='Balancing item: equals liabilities + equity less other assets')
# drivers
section(bs,DRV,'Drivers (change these on camera)',5)
put(bs,d['dd'],'Debtor days (on revenue)',['',58,62,67],font=BLUE,fmt=DAYS,fill=YEL)
put(bs,d['cd'],'Creditor days (on materials + subcontractors)',['',38,40,36],font=BLUE,fmt=DAYS,fill=YEL)
put(bs,d['capex'],'Capex: vans and equipment bought',['',n(60),n(80),n(90)],font=BLUE)
put(bs,d['div'],'Dividends paid to owner',['',n(220),n(260),n(240)],font=BLUE)

# ---------------- Cash flow (B..D = FY23..FY25 ; BS cols C..E)
cf=sheet('Cash Flow','Kestrel Plumbing & Gas Pty Ltd','Cash flow statement (indirect method) · years ended 30 June · A$',['FY23','FY24','FY25'])
C={}; r=6
bsc={'B':'C','C':'D','D':'E'}
def cl_(key,label,tmpl,bold=False,border=None,green=True,fill=None,note=None):
    global r
    vals=[tmpl.format(c=c,b=bsc[c],p=chr(ord(bsc[c])-1)) for c in 'BCD']
    put(cf,r,label,vals,bold=bold,border=border,fill=fill,note=note)
    if green and not bold:
        for c in 'BCD': cf[f'{c}{r}'].font=GRN
    C[key]=r; r+=1
BSr="'Balance Sheet'!"
section(cf,r,'Operating activities',4); r+=1
cl_('np','Net profit after tax',"='P&L'!{c}%d"%PL['npat'])
cl_('da','Add: depreciation',"='P&L'!{c}%d"%PL['da'])
for k,lab in [('deb','(Increase) in trade debtors'),('wip','(Increase) in contract assets'),('inv','(Increase) in inventory'),('pre','(Increase) in prepayments')]:
    cl_('a'+k,lab,f"=-({BSr}{{b}}{B[k]}-{BSr}{{p}}{B[k]})")
for k,lab in [('cre','Increase in trade creditors'),('accr','Increase in accrued expenses'),('gst','Increase in GST payable'),('payg','Increase in PAYG and super payable'),('al','Increase in annual leave provision'),('lslc','Increase in LSL provision (current)'),('lsln','Increase in LSL provision (non-current)'),('itp','Increase in income tax payable')]:
    cl_('l'+k,lab,f"={BSr}{{b}}{B[k]}-{BSr}{{p}}{B[k]}")
cl_('cfo','Net cash from operating activities','=SUM({c}%d:{c}%d)'%(C['np'],C['litp']),bold=True,border=TOP,green=False)
r+=1; section(cf,r,'Investing activities',4); r+=1
cl_('capex','Purchase of vans and equipment',f"=-{BSr}{{b}}{d['capex']}")
cl_('cfi','Net cash from investing activities','={c}%d'%C['capex'],bold=True,border=TOP,green=False)
r+=1; section(cf,r,'Financing activities',4); r+=1
cl_('ef','Net equipment finance drawn / (repaid)',f"=({BSr}{{b}}{B['efc']}+{BSr}{{b}}{B['efn']})-({BSr}{{p}}{B['efc']}+{BSr}{{p}}{B['efn']})")
cl_('dl','Loans to director',f"=-({BSr}{{b}}{B['dl']}-{BSr}{{p}}{B['dl']})",note='Cash out to the owner, outside the P&L')
cl_('dv','Dividends paid',f"=-{BSr}{{b}}{d['div']}")
cl_('cff','Net cash from financing activities','=SUM({c}%d:{c}%d)'%(C['ef'],C['dv']),bold=True,border=TOP,green=False)
r+=1; cl_('net','Net change in cash','={c}%d+{c}%d+{c}%d'%(C['cfo'],C['cfi'],C['cff']),bold=True,border=TOP,green=False)
cl_('open','Opening cash',f"={BSr}{{p}}{B['cash']}")
cl_('close','Closing cash','={c}%d+{c}%d'%(C['net'],C['open']),bold=True,border=DBL,green=False)
cl_('chk','Check: closing cash less balance sheet cash (must be 0)',f"=ROUND({{c}}{C['close']}-{BSr}{{b}}{B['cash']},0)",fill=RED,green=False)
r+=1
cl_('fcf','Free cash flow (operating less capex)','={c}%d+{c}%d'%(C['cfo'],C['cfi']),bold=True,green=False)
cl_('conv','Cash conversion (operating cash / EBITDA)',"=IF('P&L'!{c}%d=0,0,{c}%d/'P&L'!{c}%d)"%(PL['ebitda'],C['cfo'],PL['ebitda']),green=False)
for c in 'BCD': cf[f'{c}{C["conv"]}'].number_format=PCT
cl_('owner','Total cash to owner (dividends + director loan)','=-({c}%d+{c}%d)'%(C['dl'],C['dv']),green=False)

# ---------------- Normalisation (FY25)
nm=sheet('Normalisation','Kestrel Plumbing & Gas Pty Ltd','EBITDA normalisation · FY25 · what the seller claims vs what a buyer accepts',['Seller claims','Buyer accepts','Gap'])
nm.column_dimensions['E'].width=70; nm.cell(5,5,'Why the buyer disagrees').font=HDR; nm.cell(5,5).fill=HFILL
r=6
put(nm,r,'Reported EBITDA FY25',["='P&L'!D%d"%PL['ebitda'],"='P&L'!D%d"%PL['ebitda'],'=B%d-C%d'%(r,r)],bold=True)
for c in 'BC': nm[f'{c}{r}'].font=Font(name=F,size=10,color='008000',bold=True)
r0=r; r+=2; section(nm,r,'Adjustments',4); r+=1; a0=r
adj=[('Owner salary',"='P&L'!D%d"%PL['os'],n(20),'A replacement general manager costs about $160K. Only the excess over market is an add-back, not the whole salary.'),
('Spouse wages',"='P&L'!D%d"%PL['ow'],0,'She does the books. Replacing her costs about $75K, so nothing comes back.'),
('Rent above market (related party)',n(24),n(24),'Market rent for the depot is about $72K. Accepted only if a new lease at market is signed at completion.'),
('Owner and spouse cars',"='P&L'!D%d"%PL['ov'],n(26),'The ute is used for site visits; about $12K stays as a real cost.'),
('One-off legal dispute',n(30),n(30),'Genuinely one-off: the dispute is settled. Accepted with the settlement deed.'),
('Bad debt spike',n(22),0,'The builder who didn\'t pay is still 21% of revenue. Risk of a repeat.'),
('Price rise from 1 July (pro forma)',n(45),0,'Not earned yet. Buyers pay for history, not a promise.'),
('Owner on the tools 2 days a week',0,-n(45),'When the owner leaves, a technician must be hired. This adjustment goes the other way.')]
for lab,s,b,why in adj:
    put(nm,r,lab,[s,b,'=B%d-C%d'%(r,r)],font=BLUE)
    for c,v in (('B',s),('C',b)):
        if isinstance(v,str): nm[f'{c}{r}'].font=GRN
    nm[f'D{r}'].font=BLK; nm.cell(r,5,why).font=BLK; nm.cell(r,5).alignment=Alignment(wrap_text=True,vertical='top'); r+=1
a1=r-1
put(nm,r,'Total adjustments',['=SUM(B%d:B%d)'%(a0,a1),'=SUM(C%d:C%d)'%(a0,a1),'=B%d-C%d'%(r,r)],bold=True,border=TOP); rt=r; r+=2
put(nm,r,'Adjusted EBITDA FY25',['=B%d+B%d'%(r0,rt),'=C%d+C%d'%(r0,rt),'=B%d-C%d'%(r,r)],bold=True,border=DBL,fill=YEL); NADJ=r; r+=2
put(nm,r,'Multiple',[4.0,4.0],font=BLUE,fmt=X,fill=YEL); NMULT=r; r+=1
put(nm,r,'Implied enterprise value',['=B%d*B%d'%(NADJ,NMULT),'=C%d*C%d'%(NADJ,NMULT),'=B%d-C%d'%(r,r)],bold=True,border=DBL); NEV=r; r+=1
nm.cell(r+1,1,'Same multiple, different EBITDA. The argument is never about the multiple first; it is about which add-backs are real.').font=IT

# ---------------- Working capital
wc=sheet('Working Capital','Kestrel Plumbing & Gas Pty Ltd','Net working capital · as at 30 June · A$',['FY22','FY23','FY24','FY25'])
W={}; r=6
def wl(key,label,tmpl,first=True,fmt=USD,bold=False,border=None,fill=None,green=True,note=None):
    global r
    cols='BCDE' if first else 'CDE'; start=2 if first else 3
    vals=[tmpl.format(c=c,l=plc.get(c,'')) for c in cols]
    put(wc,r,label,vals,fmt=fmt,bold=bold,border=border,fill=fill,start=start,note=note)
    if green and not bold:
        for c in cols: wc[f'{c}{r}'].font=GRN
    W[key]=r; r+=1
section(wc,r,'Operating working capital',5); r+=1
for k,lab in [('deb','Trade debtors'),('wip','Contract assets'),('inv','Inventory'),('pre','Prepayments')]: wl(k,lab,f"={BSr}{{c}}{B[k]}")
for k,lab in [('cre','Less trade creditors'),('accr','Less accrued expenses'),('gst','Less GST payable'),('payg','Less PAYG and super payable')]: wl(k,lab,f"=-{BSr}{{c}}{B[k]}")
wl('nwc','Net working capital','=SUM({c}%d:{c}%d)'%(W['deb'],W['payg']),bold=True,border=TOP,green=False)
wl('pct','NWC as % of revenue',"=IF('P&L'!{l}%d=0,0,{c}%d/'P&L'!{l}%d)"%(PL['rev'],W['nwc'],PL['rev']),first=False,fmt=PCT,green=False)
wc.cell(r,1,'Excluded on purpose: cash, director loan, tax payable, leave provisions, equipment finance (cash or debt-like, see Valuation Bridge).').font=IT; r+=2
section(wc,r,'Days',5); r+=1
wl('dso','Debtor days',"={c}%d/'P&L'!{l}%d*365"%(W['deb'],PL['rev']),first=False,fmt=DAYS,green=False)
wl('dwip','WIP days (on revenue)',"={c}%d/'P&L'!{l}%d*365"%(W['wip'],PL['rev']),first=False,fmt=DAYS,green=False)
wl('dio','Stock days (on materials)',"={c}%d/'P&L'!{l}%d*365"%(W['inv'],PL['cm']),first=False,fmt=DAYS,green=False)
wl('dpo','Creditor days (on materials + subcontractors)',"=-{c}%d/('P&L'!{l}%d+'P&L'!{l}%d)*365"%(W['cre'],PL['cm'],PL['cs']),first=False,fmt=DAYS,green=False)
wl('ccc','Cash conversion cycle','={c}%d+{c}%d+{c}%d-{c}%d'%(W['dso'],W['dwip'],W['dio'],W['dpo']),first=False,fmt=DAYS,bold=True,border=TOP,green=False)
r+=1; section(wc,r,'NWC peg and completion adjustment',5); r+=1
put(wc,r,'Peg: average NWC % of revenue, FY23–FY25',['','','','=AVERAGE(C%d:E%d)'%(W['pct'],W['pct'])],fmt=PCT); W['pegp']=r; r+=1
put(wc,r,'Peg in dollars (on FY25 revenue)',['','','',"=E%d*'P&L'!D%d"%(W['pegp'],PL['rev'])],bold=True,fill=YEL); W['peg']=r; r+=1
put(wc,r,'NWC delivered at completion (FY25 close)',['','','','=E%d'%W['nwc']]); W['act']=r; r+=1
put(wc,r,'Completion adjustment (+ buyer pays more / − price falls)',['','','','=E%d-E%d'%(W['act'],W['peg'])],bold=True,border=DBL); W['adj']=r; r+=1
wc.cell(r+1,1,'Every extra debtor day at FY25 revenue ties up about this much cash:').font=IT
wc.cell(r+1,5,"='P&L'!D%d/365"%PL['rev']).number_format=USD

# ---------------- Valuation bridge
vb=sheet('Valuation Bridge','Kestrel Plumbing & Gas Pty Ltd','Enterprise value to equity value · cash-free, debt-free · A$',['Buyer view','Seller view'])
vb.column_dimensions['D'].width=70; vb.cell(5,4,'Note').font=HDR; vb.cell(5,4).fill=HFILL
r=6; V={}
def vl(key,label,b,s,bold=False,border=None,fill=None,note=None,fmt=USD,font=BLK):
    global r; put(vb,r,label,[b,s],bold=bold,border=border,fill=fill,fmt=fmt,font=font)
    if note: vb.cell(r,4,note).font=IT
    V[key]=r; r+=1
E='E'
vl('ebitda','Adjusted EBITDA FY25',"=Normalisation!C%d"%NADJ,"=Normalisation!B%d"%NADJ)
vl('mult','Multiple',"=Normalisation!C%d"%NMULT,"=Normalisation!B%d"%NMULT,fmt=X)
vl('ev','Enterprise value','=B%d*B%d'%(V['ebitda'],V['mult']),'=C%d*C%d'%(V['ebitda'],V['mult']),bold=True,border=TOP)
for c in 'BC':
    for k in ('ebitda','mult'): vb[f'{c}{V[k]}'].font=GRN
r+=1; section(vb,r,'Less debt and debt-like items',3); r+=1
vl('ef','Equipment finance',f"=-({BSr}E{B['efc']}+{BSr}E{B['efn']})",f"=-({BSr}E{B['efc']}+{BSr}E{B['efn']})",note='Real debt. Always comes off.')
vl('lsl','Long service leave owed to staff',f"=-({BSr}E{B['lslc']}+{BSr}E{B['lsln']})",0,note='Seller says it\'s part of the business. Buyer says it\'s a debt the seller built up.')
vl('al','Annual leave owed to staff',f"=-{BSr}E{B['al']}",0,note='Often negotiated: sometimes in NWC, sometimes debt-like.')
vl('tax','Income tax payable',f"=-{BSr}E{B['itp']}",f"=-{BSr}E{B['itp']}",note='Tax on profits the seller already earned.')
vl('dl','Director loan settled before completion',0,0,note='Excluded: the $140K Div 7A loan must be repaid or cleared by the seller, not bought.')
for c in 'BC':
    for k in ('ef','lsl','al','tax'):
        if isinstance(vb[f'{c}{V[k]}'].value,str): vb[f'{c}{V[k]}'].font=GRN
vl('debt','Total debt and debt-like','=SUM(B%d:B%d)'%(V['ef'],V['dl']),'=SUM(C%d:C%d)'%(V['ef'],V['dl']),bold=True,border=TOP)
r+=1; section(vb,r,'Plus cash and working capital',3); r+=1
vl('mincash','Minimum operating cash left in the business',100000,100000,font=BLUE,note='Cash needed to run payroll; not a bonus for the seller.')
vl('scash','Surplus cash paid to seller',f"=MAX(0,{BSr}E{B['cash']}-B%d)"%(r-1),f"=MAX(0,{BSr}E{B['cash']}-C%d)"%(r-1))
vl('nwc','NWC adjustment vs peg',"='Working Capital'!E%d"%W['adj'],0,note='Seller ignores the peg. Buyer applies it.')
r+=1; vl('eq','Equity value (price for the shares)','=B%d+B%d+B%d+B%d'%(V['ev'],V['debt'],V['scash'],V['nwc']),'=C%d+C%d+C%d+C%d'%(V['ev'],V['debt'],V['scash'],V['nwc']),bold=True,border=DBL,fill=YEL)
vl('gap','Gap between the two views','=C%d-B%d'%(V['eq'],V['eq']),'',bold=True)

# ---------------- Funding & DSCR
fd=sheet('Funding & DSCR','Kestrel Plumbing & Gas Pty Ltd','How the deal is paid for, and whether the business can carry it · A$',['Seller asking','Buyer price','Buyer price, restructured'])
fd.column_dimensions['E'].width=60; fd.column_dimensions['D'].width=24; r=6; D={}
fd.cell(5,5,'Note').font=HDR; fd.cell(5,5).fill=HFILL
def fl(key,label,vals,bold=False,border=None,fill=None,fmt=USD,font=BLK,note=None):
    global r
    if isinstance(vals,str): vals=[vals.format(c=c) for c in 'BCD']
    put(fd,r,label,vals,bold=bold,border=border,fill=fill,fmt=fmt,font=font)
    for i,v in enumerate(vals):
        if isinstance(v,str) and '!' in v: fd.cell(r,2+i).font=Font(name=F,size=10,color='008000',bold=bold)
    if note: fd.cell(r,5,note).font=IT
    D[key]=r; r+=1
fl('price','Enterprise value (purchase price)',["='Valuation Bridge'!C%d"%V['ev'],"='Valuation Bridge'!B%d"%V['ev'],"='Valuation Bridge'!B%d"%V['ev']],note='Column 3 keeps the buyer price and changes only the terms.')
r+=1; section(fd,r,'Funding stack',4); r+=1
fl('bp','Bank share of price',[0.6,0.6,0.4],fmt=PCT,font=BLUE,fill=YEL)
fl('vp','Vendor finance share of price',[0.4,0.4,0.6],fmt=PCT,font=BLUE,fill=YEL)
fl('ep','Buyer cash','=1-{c}%d-{c}%d'%(D['bp'],D['vp']),fmt=PCT,note='No money down when this is 0%')
fl('bank','Bank loan','={c}%d*{c}%d'%(D['price'],D['bp']))
fl('vend','Vendor loan','={c}%d*{c}%d'%(D['price'],D['vp']))
r+=1; section(fd,r,'Terms',4); r+=1
fl('br','Bank rate',[0.085,0.085,0.085],fmt=PCT,font=BLUE)
fl('bt','Bank term (years)',[5,5,10],fmt='0',font=BLUE,fill=YEL)
fl('vr','Vendor rate',[0.04,0.04,0.04],fmt=PCT,font=BLUE)
fl('vt','Vendor term (years)',[7,7,15],fmt='0',font=BLUE,fill=YEL,note='Long vendor terms are what make the deal carry.')
fl('bpay','Bank repayment per year (P&I)','=PMT({c}%d,{c}%d,-{c}%d)'%(D['br'],D['bt'],D['bank']))
fl('vpay','Vendor repayment per year (P&I)','=PMT({c}%d,{c}%d,-{c}%d)'%(D['vr'],D['vt'],D['vend']))
fl('ds','Total debt service','={c}%d+{c}%d'%(D['bpay'],D['vpay']),bold=True,border=TOP)
r+=1; section(fd,r,'Cash available for debt service (year 1)',4); r+=1
fl('e','Adjusted EBITDA (buyer view in every column)',"='Valuation Bridge'!B%d"%V['ebitda'],note='The bank lends on real EBITDA, whatever price you agree.')
fl('mc','Less maintenance capex (3-yr average)',f"=-ROUND(AVERAGE({BSr}C{d['capex']}:E{d['capex']}),0)")
fl('ti','Less tax (EBITDA less D&A less year-1 interest, x rate)',"=-ROUND(MAX(0,({c}%d-'P&L'!D%d-{c}%d*{c}%d-{c}%d*{c}%d)*'P&L'!$B$%d),0)"%(D['e'],PL['da'],D['bank'],D['br'],D['vend'],D['vr'],PL['tr']))
fl('cfads','CFADS','=SUM({c}%d:{c}%d)'%(D['e'],D['ti']),bold=True,border=TOP)
r+=1; fl('dscr','DSCR (CFADS / debt service)','=IF({c}%d=0,0,{c}%d/{c}%d)'%(D['ds'],D['cfads'],D['ds']),bold=True,fmt=X,fill=YEL,border=DBL)
fl('tgt','Target DSCR',[1.5,1.5,1.5],fmt=X,font=BLUE)
fl('ok','Verdict','=IF({c}%d>={c}%d,"Fundable","Fails")'%(D['dscr'],D['tgt']),bold=True,fill=RED)
r+=1; section(fd,r,'What the business can carry',4); r+=1
fl('maxds','Max debt service at target DSCR','={c}%d/{c}%d'%(D['cfads'],D['tgt']))
fl('room','Headroom / (shortfall) per year','={c}%d-{c}%d'%(D['maxds'],D['ds']),bold=True)

# ---------------- Red flags
rf=sheet('Red Flags','Kestrel Plumbing & Gas Pty Ltd','Diligence checks calculated from the statements',['FY23','FY24','FY25'])
rf.column_dimensions['E'].width=70; rf.cell(5,5,'Why it matters').font=HDR; rf.cell(5,5).fill=HFILL
r=6
def rl(label,tmpl,fmt,why,bold=False):
    global r
    put(rf,r,label,[tmpl.format(c=c,b=bsc[c]) for c in 'BCD'],fmt=fmt,bold=bold)
    rf.cell(r,5,why).font=BLK; rf.cell(r,5).alignment=Alignment(wrap_text=True,vertical='top'); r+=1
put(rf,r,'Largest customer revenue (one builder)',[n(640),n(780),n(900)],font=BLUE); top=r
rf.cell(r,5,'Input: from the aged debtors and sales ledger.').font=IT; r+=1
rl('Largest customer % of revenue',"={c}%d/'P&L'!{c}%d"%(top,PL['rev']),PCT,'Over 20% and rising. Lose the builder and about a fifth of revenue goes with it.')
put(rf,r,'Revenue growth',['',"='P&L'!C%d/'P&L'!B%d-1"%(PL['rev'],PL['rev']),"='P&L'!D%d/'P&L'!C%d-1"%(PL['rev'],PL['rev'])],fmt=PCT)
rf.cell(r,5,'Healthy top line. The question is what kind of revenue.').font=BLK; r+=1
rl('Gross margin',"='P&L'!{c}%d"%PL['gm'],PCT,'Falling every year while projects grow. Projects work for the builder is probably low margin.')
rl('Debtor days',"='Working Capital'!{b}%d"%W['dso'],DAYS,'Stretching. Customers are paying slower, and the builder is the likely cause.')
rl('Bad debts % of revenue',"='P&L'!{c}%d/'P&L'!{c}%d"%(PL['ob'],PL['rev']),PCT,'Tripled in FY25. Same customer as the legal dispute.')
rl('Cash conversion (operating cash / EBITDA)',"='Cash Flow'!{c}%d"%C['conv'],PCT,'Below 80% means profit is turning into debtors, not cash.')
rl('Cash taken by owner (dividends + director loan)',"='Cash Flow'!{c}%d"%C['owner'],USD,'Compare with free cash flow. Is the owner taking more than the business makes?')
rl('Free cash flow',"='Cash Flow'!{c}%d"%C['fcf'],USD,'')
rl('Loan to director (Div 7A)',"='Balance Sheet'!{b}%d"%B['dl'],USD,'Cash leaving the business without going through the P&L. Must be cleared before completion.')
rl('Staff leave owed (annual + long service)',"='Balance Sheet'!{b}%d+'Balance Sheet'!{b}%d+'Balance Sheet'!{b}%d"%(B['al'],B['lslc'],B['lsln']),USD,'Unfunded promise to staff. Who pays it is a price negotiation.')
rl('Related-party rent',"='P&L'!{c}%d"%PL['ore'],USD,'Paid to the owner\'s trust, above market. Needs a new lease at market rent at completion.')
rl('Owner salary as % of EBITDA',"='P&L'!{c}%d/'P&L'!{c}%d"%(PL['os'],PL['ebitda']),PCT,'And the owner is on the tools two days a week. A buyer has to replace both jobs.')
for row in rf.iter_rows(min_row=7,max_row=r-1,min_col=2,max_col=4):
    for c in row:
        if isinstance(c.value,str) and c.value.startswith('='): c.font=GRN

# ---------------- Video ideas
vi=wb.create_sheet('Video ideas'); vi.sheet_view.showGridLines=False
vi.column_dimensions['A'].width=78; vi.column_dimensions['B'].width=20; vi.column_dimensions['C'].width=70
vi['A1']='Technical videos from this one target'; vi['A1'].font=H1
for i,h in enumerate(['Title','Tab','What to show on screen']): c=vi.cell(3,i+1,h); c.font=HDR; c.fill=HFILL
ideas=[
('The Seller Said $968K EBITDA. I Got $614K. (Add-Back Teardown)','Normalisation','Go line by line: which add-backs survive and why. End on the $1.4M valuation gap.'),
('If You Can\'t Read A Balance Sheet, You Can\'t Buy A Business','Balance Sheet','Walk top to bottom: what each line means for a buyer, then the 3 lines that change the price.'),
('Enterprise Value vs Equity Value: Where $500K Disappears At Settlement','Valuation Bridge','Same business, two bridges. Debt-like items, leave provisions, surplus cash.'),
('The Most Valuable Working Capital Training You\'ll Ever Watch (The NWC Peg)','Working Capital','Build the peg live, then change debtor days by 10 and show what it costs at completion.'),
('Every Line Of A P&L A Buyer Questions (Real Example)','P&L','Owner salary, family wages, related-party rent, bad debts, legal. What each hides.'),
('Profit Isn\'t Cash: Reading A Cash Flow Statement Like A Buyer','Cash Flow','Profit vs operating cash, cash conversion, and cash quietly taken by the owner.'),
('DSCR Explained With A Real Deal: Same Price, Different Terms, Different Answer','Funding & DSCR','Three columns: seller price fails, buyer price on standard terms fails, buyer price on restructured terms passes 1.5×.'),
('11 Red Flags Hiding In A $4.2M Business\'s Financials','Red Flags','One by one, each calculated from the statements.'),
('What Is A Div 7A Loan, And Why It Matters When You Buy','Balance Sheet','The $140K loan to the director, and how to clear it before completion.'),
('Annual Leave, Long Service Leave: The $230K Debt Nobody Puts In The Price','Valuation Bridge','Who pays accrued leave, and how it changes the price.'),
('Watch Me Price This $4.2M Plumbing Business Live (Full Model)','All tabs','End to end in one sitting: P&L, normalise, bridge, fund, DSCR.'),
('Gross Margin Is Falling And Revenue Is Rising. Here\'s Why That\'s A Problem','P&L / Red Flags','Mix shift to projects for one builder; the margin story behind growth.'),
]
for i,(t,tab,s) in enumerate(ideas,4):
    vi.cell(i,1,t).font=BLK; vi.cell(i,2,tab).font=BLK; vi.cell(i,3,s).font=BLK
    vi.cell(i,3).alignment=Alignment(wrap_text=True,vertical='top'); vi.cell(i,1).alignment=Alignment(wrap_text=True,vertical='top')
for ws in wb.worksheets: ws.sheet_properties.pageSetUpPr.fitToPage=True
wb.save('Kestrel Plumbing - Example Target Financials.xlsx'); print('saved')
